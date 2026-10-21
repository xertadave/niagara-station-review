"""Regression tests – one per required behaviour or fixed bug.

Run before every build:   python build.py        (or: python -m unittest discover -s tests -v)
A build is not packaged unless every test passes.  When a bug is fixed, add a test here that
fails on the old code, and a line in CHANGELOG.md.
"""
import hashlib
import json
import os
import re
import shutil
import socket
import subprocess
import sys
import tempfile
import time
import unittest
import urllib.request
import zipfile
from datetime import datetime

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import fixtures  # noqa: E402
from station_review import store, report  # noqa: E402
from station_review.pipeline import analyse  # noqa: E402
from station_review.checks import readings  # noqa: E402

TODAY = datetime(2026, 10, 20)


def read(*p):
    with open(os.path.join(*p), encoding='utf-8', errors='ignore') as fh:
        return fh.read()


def codes(items, kind=None):
    return {i['match'] for i in items if kind is None or i['kind'] == kind}


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix='nsr-test-')
        store.set_data(os.path.join(self.tmp, 'data'))
        self.bog = fixtures.build_bog()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def review(self, site='Site A'):
        F, items, summary = analyse(self.bog, TODAY)
        return store.create_review(site, F['station'], F['saved'], F['niagara_version'], self.bog, items, 'x.bog', summary)

    def item(self, rid, match):
        _, items = store.get_review(rid)
        return next(i for i in items if i['match'] == match)


# ---------------------------------------------------------------------------------------------
class TestReading(Base):
    def test_bog_with_illegal_character_reference_is_read(self):
        """v1.0.0: firmware strings can hold &#x0;, which XML forbids; the file must still be read."""
        self.assertIn('&#x0;', fixtures.build_xml())
        F, items, _ = analyse(self.bog, TODAY)
        self.assertEqual(F['station'], 'PLANT01')
        self.assertEqual(F['saved'], '2026-10-01 09:30')

    def test_non_station_file_refused(self):
        with self.assertRaises(ValueError):
            analyse(b'hello', TODAY)


class TestChecks(Base):
    def setUp(self):
        super().setUp()
        self.F, self.items, self.summary = analyse(self.bog, TODAY)
        self.found = codes(self.items, 'action')

    def test_every_check_finds_its_example(self):
        for m in ['OVR:Mode', 'FAULT:RoomTemp', 'SCHED-EMPTY:EmptySched', 'ALARM-ANALOG', 'ALARM-ROUTE', 'USERS', 'SEC', 'BACKUP',
                  'UNITS:Temperature_UI5', 'UNITS:ExtractHum', 'NAMING:VAV', 'MISSING:Damper', 'DISABLED:Damper', 'FALLBACK:Setpoint',
                  'VERSION', 'HISTORY', 'MBUS-DUP', 'FROST:FrostSP']:
            self.assertIn(m, self.found)

    def test_emergency_override_is_high(self):
        o = next(i for i in self.items if i['match'] == 'OVR:Mode')
        self.assertEqual(o['priority'], 'High')
        self.assertIn('VAV-1-04', o['title'])

    def test_wiresheet_setting_is_status_not_override(self):
        """A wiresheet flag set by hand (e.g. an enrolment switch) is a setting, not a fault."""
        self.assertNotIn('OVR:Enrol', self.found)
        self.assertIn('ST-SETTINGS', codes(self.items, 'status'))

    def test_energy_counters_are_not_sensor_faults(self):
        """v1.0.0: kWh registers read large numbers; only temperature, humidity and CO2 are judged."""
        self.assertFalse(any(m.startswith('FAULT:') and 'Energy' in m for m in self.found))

    def test_readings_one_per_input(self):
        r = readings(self.F)
        self.assertEqual(sorted(x['point'] for x in r), ['ExtractHum', 'RoomTemp', 'Supply'])
        self.assertTrue(next(x for x in r if x['point'] == 'RoomTemp')['fault'])

    def test_naming_outlier_named(self):
        n = next(i for i in self.items if i['match'] == 'NAMING:VAV')
        self.assertIn('VAV-9', n['observed'])
        self.assertEqual(n['status'], 'Client decision')

    def test_points_in_manual_counted(self):
        m = next(i for i in self.items if i['code'] == 'ST-MANUAL')
        self.assertEqual(m['count'], 1)
        self.assertEqual(m['result'], 'See actions')


class TestDecisions(Base):
    def test_site_rule_from_disregard_skips_next_time(self):
        """A Disregard with 'This site' becomes a rule; the next review at the same site has it under Already decided."""
        r1 = self.review('Site A')
        a = self.item(r1, 'BACKUP')
        store.update_item(r1, a['id'], {'decision': 'disregard', 'disregardReason': 'Backups are kept on the supervisor'})
        store.update_item(r1, a['id'], {'ruleScope': 'site'})
        self.assertEqual(len(store.rules()), 1)
        r2 = self.review('site a ')                      # same site, different spelling of case/space
        b = self.item(r2, 'BACKUP')
        self.assertEqual(b['kind'], 'precedent')
        self.assertEqual(b['decision'], 'disregard')
        r3 = self.review('Site B')                       # other site: still raised
        self.assertEqual(self.item(r3, 'BACKUP')['kind'], 'action')

    def test_all_sites_rule(self):
        r1 = self.review('Site A')
        a = self.item(r1, 'VERSION')
        store.update_item(r1, a['id'], {'decision': 'disregard', 'ruleScope': 'all', 'disregardReason': 'patch level only'})
        self.assertEqual(self.item(self.review('Site C'), 'VERSION')['kind'], 'precedent')

    def test_rule_removed_when_decision_changes(self):
        r1 = self.review()
        a = self.item(r1, 'SEC')
        store.update_item(r1, a['id'], {'decision': 'disregard', 'ruleScope': 'all'})
        store.update_item(r1, a['id'], {'decision': 'include'})
        self.assertEqual(store.rules(), [])

    def test_report_only_disregard_is_precedent_at_same_site_only(self):
        r1 = self.review('Site A')
        a = self.item(r1, 'ALARM-ANALOG')
        store.update_item(r1, a['id'], {'decision': 'disregard', 'disregardReason': 'site team adjust setpoints'})
        b = self.item(self.review('Site A'), 'ALARM-ANALOG')
        self.assertEqual(b['kind'], 'precedent')
        self.assertIn('site team adjust setpoints', b['disregardReason'])
        self.assertEqual(self.item(self.review('Site B'), 'ALARM-ANALOG')['kind'], 'action')

    def test_included_finding_carried_but_generated_wording_rebuilt(self):
        """v1.0.0: only wording the reviewer wrote is carried; generated wording follows the new station's data."""
        r1 = self.review()
        u = self.item(r1, 'USERS')
        store.update_item(r1, u['id'], {'decision': 'include'})
        n = self.item(r1, 'NAMING:VAV')
        store.update_item(r1, n['id'], {'decision': 'include', 'clientText': 'My own wording'})
        r2 = self.review()
        u2, n2 = self.item(r2, 'USERS'), self.item(r2, 'NAMING:VAV')
        self.assertEqual(u2['carried'], 'PLANT01')
        self.assertEqual(u2['decision'], 'undecided')            # the reviewer still confirms it
        self.assertNotIn('wordingCarried', u2)
        self.assertEqual(n2['clientText'], 'My own wording')
        self.assertTrue(n2['wordingCarried'])

    def test_rule_applies_patterns(self):
        it = {'code': 'MISSING', 'match': 'MISSING:Fan Speed'}
        r = {'active': True, 'scope': 'all', 'applies': ['DISABLED:Fan Speed', 'missing:fan speed']}
        self.assertTrue(store.rule_applies(r, it, 'x'))
        self.assertTrue(store.rule_applies({**r, 'applies': ['MISS*']}, it, 'x'))
        self.assertFalse(store.rule_applies({**r, 'active': False}, it, 'x'))


class TestReport(Base):
    def build(self, settings=None):
        rid = self.review('Example Site')
        meta, items = store.get_review(rid)
        for i in items:
            if i['kind'] != 'precedent':
                store.update_item(rid, i['id'], {'decision': 'include'})
        f = self.item(rid, 'FROST:FrostSP')
        store.update_item(rid, f['id'], {'status': 'Completed during review'})
        h = self.item(rid, 'HISTORY')
        store.update_item(rid, h['id'], {'decision': 'disregard', 'listAsReviewed': True, 'disregardReason': 'The supervisor keeps the trends.'})
        meta, items = store.get_review(rid)
        out = os.path.join(self.tmp, 'r.docx')
        s = {**store.settings(), **(settings or {})}
        report.build(meta, items, s, out, report.template_path(store.DATA, s))
        from docx import Document
        d = Document(out)
        text = '\n'.join(t.text or '' for t in d.element.body.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
        foot = '\n'.join(p.text for sec in d.sections for p in sec.footer.paragraphs)
        return text, foot

    def test_report_sections_and_no_placeholders(self):
        text, foot = self.build({'author': 'A Person', 'role': 'Engineer'})
        for s in ['Remote PPM Report', 'Executive Summary', 'Site and System Information', 'System Status', 'Readings at the Time of Review',
                  'Actions and Recommendations', 'Completed During Review', 'Reviewed – No Action Required', 'Engineer Notes', 'Contact us',
                  'Example Site', 'PLANT01', 'A Person', 'The supervisor keeps the trends.']:
            self.assertIn(s, text)
        self.assertNotIn('{{', text + foot)
        self.assertNotIn('Energy Audit Report', text)
        self.assertNotIn('Building Information', text)      # template body removed, not just the cover replaced
        self.assertIn('Example Site', foot)

    def test_contents_entry_not_taken_for_contact_block(self):
        """v1.0.0: the template's contents page lists 'Contact us'; the body was only cleared up to that entry."""
        text, _ = self.build()
        self.assertNotIn('Contents', text)                      # the contents page went with the template body
        from docx import Document
        d = Document(os.path.join(self.tmp, 'r.docx'))
        W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        starts = [el for el in d.element.body.iterchildren() if el.tag == W + 'p' and
                  ''.join(t.text or '' for t in el.iter(W + 't')).strip().startswith('Contact us')]
        self.assertEqual(len(starts), 1)


class TestServer(Base):
    def test_upload_review_and_download(self):
        s = socket.socket()
        s.bind(('127.0.0.1', 0))
        port = s.getsockname()[1]
        s.close()
        p = subprocess.Popen([sys.executable, '-m', 'station_review', '--no-browser', '--port', str(port), '--data', os.path.join(self.tmp, 'srv')],
                             cwd=ROOT, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        try:
            for _ in range(100):
                try:
                    socket.create_connection(('127.0.0.1', port), timeout=0.2).close()
                    break
                except OSError:
                    time.sleep(0.1)
            base = f'http://127.0.0.1:{port}'
            op = urllib.request.build_opener(urllib.request.ProxyHandler({}))
            bnd = 'XXbnd'
            body = (f'--{bnd}\r\nContent-Disposition: form-data; name="site"\r\n\r\nSite Z\r\n'
                    f'--{bnd}\r\nContent-Disposition: form-data; name="file"; filename="plant.bog"\r\nContent-Type: application/octet-stream\r\n\r\n').encode() + \
                self.bog + f'\r\n--{bnd}--\r\n'.encode()
            r = op.open(urllib.request.Request(base + '/api/upload', data=body, headers={'Content-Type': f'multipart/form-data; boundary={bnd}'}))
            rid = json.load(r)['id']
            j = json.load(op.open(base + '/api/review/' + rid))
            first = j['items'][0]
            op.open(urllib.request.Request(f'{base}/api/review/{rid}/item/{first["id"]}', data=json.dumps({'decision': 'include'}).encode(), method='POST'))
            op.open(urllib.request.Request(f'{base}/api/review/{rid}/report', data=b'{}', method='POST'))
            doc = op.open(f'{base}/api/review/{rid}/report.docx').read()
            self.assertEqual(doc[:2], b'PK')
            # v1.0.0: the import route was taken for a rule id ('rule not found')
            rules = [{'id': 'r1', 'scope': 'site', 'site': '{{SITE}}', 'applies': ['BACKUP'], 'text': 't', 'reason': 'r'}]
            j = json.load(op.open(urllib.request.Request(f'{base}/api/rules/import', data=json.dumps({'rules': rules, 'site': 'Site Z'}).encode(), method='POST')))
            self.assertEqual(j['added'], 1)
            j = json.load(op.open(base + '/api/rules'))
            self.assertEqual(j[0]['site'], 'Site Z')
            for page in ('/', '/review', '/style.css', '/logo.png'):
                self.assertEqual(op.open(base + page).status, 200)
        finally:
            p.terminate()
            p.wait(10)


class TestPage(unittest.TestCase):
    def test_hidden_attribute_wins_over_display(self):
        """v1.0.0: panels with display:grid ignored the hidden attribute, so every card showed its 'why' panel."""
        css = read(ROOT, 'station_review', 'web', 'style.css')
        self.assertRegex(css, r'\[hidden\]\{display:none!important\}')

    def test_review_page_does_not_rebuild_cards_on_save(self):
        """Saving a field must update in place, never rebuild the list (that jumped the page back to the top)."""
        js = read(ROOT, 'station_review', 'web', 'review.html')
        save = js[js.index('async function save'):js.index('function el(')]
        self.assertNotIn('load(', save)
        self.assertNotIn('innerHTML', save)


class TestHousekeeping(unittest.TestCase):
    # Hashes (first 16 hex chars of sha256) of site, station, people and address words from the stations this
    # tool was first built against.  The words themselves are never written in the repository.
    BANNED = {'1efbff7c1cec1720', '5c0055fa837767cb', '7f11017264c937c1', '4c4581a33c876214', '84e6af8a53ad2a8e', 'e0186de91d8c03f4',
              '5913daa7d2d2e475', 'd5a8aab4d830ba51', 'aed0521ab0f2b747', '551f1957662b6f98', '4e93cd22c2ea1016', 'f5d22ddd4aa36114',
              '09dcc7a7d81227b6', '4dba4870ec1c809b', 'b398a038f25ea585', 'e1b683e26a3aad21', '382ec6cf542b3df1', 'ec926e3d021087fe',
              'fc66ac973774d53b', '3c70d62a991c6b60', 'ef3b71dc955ecaa4'}

    def _texts(self):
        for d, dirs, files in os.walk(ROOT):
            dirs[:] = [x for x in dirs if x not in ('.git', '__pycache__', 'data', 'dist')]
            for f in files:
                fp = os.path.join(d, f)
                if f.endswith('.docx'):
                    with zipfile.ZipFile(fp) as z:
                        yield fp, ' '.join(z.read(n).decode('utf-8', 'ignore') for n in z.namelist() if n.endswith('.xml'))
                elif f.endswith(('.py', '.html', '.js', '.css', '.md', '.txt', '.bat', '.json', '.cfg', '.yaml', '.gitignore')):
                    yield fp, read(fp)

    def test_no_site_references_anywhere(self):
        """v1.0.0: nothing in the repository may name a site, station, person or address from a real job."""
        bad = []
        for fp, t in self._texts():
            for w in set(re.findall(r'[A-Za-z0-9_.\-]+', t.lower())):
                for cand in {w, w.strip('.-_')}:
                    if hashlib.sha256(cand.encode()).hexdigest()[:16] in self.BANNED:
                        bad.append(os.path.relpath(fp, ROOT))
        self.assertEqual(sorted(set(bad)), [])

    def test_every_import_is_in_requirements(self):
        import ast
        std = set(sys.stdlib_module_names)
        req = {l.split('>=')[0].split('==')[0].strip().lower().replace('-', '_') for l in read(ROOT, 'requirements.txt').splitlines() if l.strip() and not l.startswith('#')}
        alias = {'docx': 'python_docx'}
        missing = set()
        for d, dirs, files in os.walk(os.path.join(ROOT, 'station_review')):
            for f in files:
                if f.endswith('.py'):
                    for node in ast.walk(ast.parse(read(d, f))):
                        names = [a.name for a in node.names] if isinstance(node, ast.Import) else \
                            [node.module] if isinstance(node, ast.ImportFrom) and node.level == 0 and node.module else []
                        for n in names:
                            top = n.split('.')[0]
                            if top not in std and alias.get(top, top) not in req:
                                missing.add(top)
        self.assertEqual(missing, set())

    def test_bat_files_have_windows_line_endings(self):
        for f in os.listdir(ROOT):
            if f.endswith('.bat'):
                b = open(os.path.join(ROOT, f), 'rb').read()
                self.assertNotIn(b'\n', b.replace(b'\r\n', b''), f)

    def test_version_has_changelog_entry(self):
        from station_review.server import VERSION
        self.assertIn(f'## v{VERSION}', read(ROOT, 'CHANGELOG.md'))


if __name__ == '__main__':
    unittest.main()
