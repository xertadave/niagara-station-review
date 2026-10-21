"""Build the Word report from a review's decisions.

Uses a Word template for the cover, headers, footers, styles and Contact us page: the user's
own template (data/branding/template.docx) if one is uploaded, otherwise templates/report_template.docx.
In the template, {{SITE_NAME}} and {{CLIENT}} are replaced, the cover title is replaced with the
report title, and the body between the cover and the Contact us block is replaced.
"""
import os

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
BUILTIN_TEMPLATE = os.path.join(os.path.dirname(HERE), 'templates', 'report_template.docx')
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

PRI = {'High': 'B42318', 'Medium': 'B25E09', 'Low': '3F6C88', 'Review': '6B4FB8'}
GREEN, NAVY, TEAL, AMBER = '24734A', '002D72', '0097A9', 'B25E09'
DONE = ('Done', 'Completed during review')
TEMPLATE_TITLE = 'Energy Audit Report'          # the cover title text in the bundled template


def _fmt(v):
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(v)
    return f'{f:.1f}'


class Writer:
    def __init__(self, template):
        self.doc = Document(template)
        body = self.doc.element.body
        kids = list(body.iterchildren())
        self.sect = kids[-1] if kids and kids[-1].tag == W + 'sectPr' else None
        cover_end = 0
        for i, el in enumerate(kids):
            if el.tag == W + 'sdt':
                cover_end = i + 1
                break
        self.contact = None
        for el in reversed(kids[cover_end:]):      # the last "Contact us" paragraph; the contents page lists one too
            txt = ''.join(t.text or '' for t in el.iter(W + 't'))
            sty = el.find('.//' + W + 'pStyle')
            if el.tag == W + 'p' and txt.strip().lower().startswith('contact us') and not (sty is not None and sty.get(W + 'val', '').upper().startswith('TOC')):
                self.contact = el
                break
        for el in kids[cover_end:]:
            if el is self.contact or el is self.sect:
                continue
            if self.contact is not None and kids.index(el) > kids.index(self.contact):
                continue
            body.remove(el)
        self.anchor = self.contact if self.contact is not None else self.sect
        if self.contact is not None:            # keep the Contact us page to one page (large closing picture)
            for inl in self.contact.iter():
                if inl.tag.split('}')[-1] != 'inline':     # only the in-line picture; anchored page artwork is left alone
                    continue
                for el in inl.iter():
                    if el.tag.split('}')[-1] in ('extent', 'ext') and el.get('cy') and int(el.get('cy')) > 5000000:
                        el.set('cx', str(int(int(el.get('cx')) * 0.8)))
                        el.set('cy', str(int(int(el.get('cy')) * 0.8)))

    def _put(self, el):
        if self.anchor is not None:
            self.anchor.addprevious(el)
        else:
            self.doc.element.body.append(el)

    def style(self, *names):
        have = {s.name for s in self.doc.styles}
        for n in names:
            if n in have:
                return n
        return None

    def para(self, text='', style=None, bold=False, keep=False):
        p = self.doc.add_paragraph(style=self.style(style, 'Body Text', 'Normal') if style else self.style('Body Text', 'Normal'))
        if text:
            r = p.add_run(text)
            r.bold = bold
        if keep:
            p.paragraph_format.keep_with_next = True
        self._put(p._p)
        return p

    def h1(self, text, numbered=True):
        return self.para(text, self.style('Numbered Heading 1' if numbered else 'Heading 1', 'Heading 1'), keep=True)

    def h2(self, text):
        return self.para(text, self.style('Numbered Heading 2', 'Heading 2'), keep=True)

    def bullet(self, text):
        return self.para(text, self.style('Bullet 1', 'List Bullet', 'Body Text'))

    def pagebreak(self):
        p = self.para('')
        p.add_run().add_break(WD_BREAK.PAGE)

    @staticmethod
    def shade(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement('w:shd')
        sh.set(qn('w:val'), 'clear')
        sh.set(qn('w:color'), 'auto')
        sh.set(qn('w:fill'), fill)
        tcPr.append(sh)

    def table(self, headers, rows, widths, head_fill=NAVY, big=None, keep=None):
        t = self.doc.add_table(rows=1, cols=len(headers))
        if self.style('Table Grid'):
            t.style = self.doc.styles['Table Grid']
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            c.text = ''
            r = c.paragraphs[0].add_run(h)
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(9.5)
            self.shade(c, head_fill)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ''
                para = cells[i].paragraphs[0]
                if isinstance(v, tuple):            # (text, colour, bold)
                    r = para.add_run(v[0])
                    r.bold = v[2]
                    r.font.color.rgb = RGBColor.from_string(v[1])
                else:
                    r = para.add_run('' if v is None else str(v))
                r.font.size = Pt(big or 9.5)
        tblPr = t._tbl.tblPr
        b = OxmlElement('w:tblBorders')
        for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            x = OxmlElement('w:' + e)
            x.set(qn('w:val'), 'single')
            x.set(qn('w:sz'), '4')
            x.set(qn('w:color'), 'BFC7CF')
            b.append(x)
        tblPr.append(b)
        lay = OxmlElement('w:tblLayout')
        lay.set(qn('w:type'), 'fixed')
        tblPr.append(lay)
        for old in tblPr.findall(qn('w:tblW')):
            tblPr.remove(old)
        tw = OxmlElement('w:tblW')
        tw.set(qn('w:w'), str(int(sum(widths) * 567)))
        tw.set(qn('w:type'), 'dxa')
        tblPr.append(tw)
        for i, gc in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))):
            gc.set(qn('w:w'), str(int(widths[i] * 567)))
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
        if keep is None:
            keep = len(rows) <= 14
        if keep:                                   # small tables stay on one page with their heading
            for row in t.rows[:-1]:
                for c in row.cells:
                    for p in c.paragraphs:
                        p.paragraph_format.keep_with_next = True
        self._put(t._tbl)
        self.para('')
        return t

    def replace_text(self, mapping, where='body'):
        if where == 'body':
            parts = [self.doc.element.body]
        else:
            parts = [x._element for s in self.doc.sections for x in (s.footer, s.header, s.first_page_footer, s.first_page_header, s.even_page_footer, s.even_page_header)]
        for root in parts:
            for t in root.iter(W + 't'):
                if t.text:
                    for a, b in mapping.items():
                        if a in t.text:
                            t.text = t.text.replace(a, b)

    def set_contact(self, author, role):
        """Bundled template's Contact us block: name runs, role runs, then a second contact that is removed.
        Only touched when the runs are laid out as expected; a user's own template is left alone."""
        if self.contact is None:
            return
        runs = [t for t in self.contact.iter(W + 't')]
        tx = [t.text or '' for t in runs]
        try:
            c = tx.index('Contact us')
        except ValueError:
            return
        if len(tx) < c + 11 or tx[c + 4].strip() != 'BMS' or tx[c + 8].strip() != 'Energy':
            return
        for k in range(c + 6, c + 11):          # second contact
            runs[k].text = ''
        if author:
            runs[c + 1].text, runs[c + 2].text, runs[c + 3].text = author, '', ''
        if role:
            runs[c + 4].text, runs[c + 5].text = role, ''

    def save(self, path):
        self.doc.save(path)


def template_path(data_dir, settings):
    own = os.path.join(data_dir, 'branding', 'template.docx')
    if settings.get('use_template', True) and os.path.exists(own):
        return own
    return BUILTIN_TEMPLATE


def build(meta, items, settings, out_path, template):
    w = Writer(template)
    site, station = meta['site'], meta['station']
    title = settings.get('report_title') or 'Remote PPM Report'
    w.replace_text({'{{SITE_NAME}}': station, '{{CLIENT}}': site, TEMPLATE_TITLE: title}, 'body')
    w.replace_text({'{{SITE_NAME}}': f"{site} – {station}", '{{CLIENT}}': site}, 'headers')
    w.set_contact(settings.get('author', '').strip(), settings.get('role', '').strip())
    if w.sect is not None:
        w.doc.core_properties.title = f"{title} – {site} {station}"

    inc = [i for i in items if i.get('decision') == 'include']
    acts = [i for i in inc if i.get('was_kind', i['kind']) == 'action']
    stat = [i for i in inc if i.get('was_kind', i['kind']) == 'status']
    order = {'High': 0, 'Medium': 1, 'Low': 2, 'Review': 3}
    acts.sort(key=lambda i: (order.get(i.get('priority'), 9), i.get('order', 0)))
    done = [a for a in acts if a.get('status') in DONE]
    client = [a for a in acts if a.get('status') == 'Client decision']
    open_ = [a for a in acts if a not in done and a not in client]
    manual = next((i for i in items if i['code'] == 'ST-MANUAL'), {}).get('count', 0)
    reviewed = [i for i in items if i.get('decision') == 'disregard' and i.get('listAsReviewed')]
    readings_item = next((i for i in stat if i['code'] == 'ST-READINGS'), None)
    stat = [s for s in stat if s['code'] != 'ST-READINGS']
    summ = meta.get('summary', {})

    # ---- executive summary
    w.h1('Executive Summary', numbered=False)
    w.para(f"A remote planned preventive maintenance (PPM) review of the BMS controller {station} at {site} was carried out by the "
           f"{settings.get('company_team') or 'Smart Buildings'} team. The review used the controller's station configuration backup, "
           f"saved {meta.get('saved') or 'for this review'}, and covered the plant, terminal units, field devices and metering connected to this controller.")
    if meta.get('summary_note'):
        for line in meta['summary_note'].split('\n'):
            if line.strip():
                w.para(line.strip())
    w.table(['Actions open', 'Completed during review', 'For client decision', 'Points in manual'],
            [[(str(len(open_)), AMBER, True), (str(len(done)), GREEN, True), (str(len(client)), '3F6C88', True),
              (str(manual), 'B42318' if manual else GREEN, True)]], [4, 4.6, 4, 3.4], head_fill=TEAL, big=20)
    w.para('The following report covers:')
    for b in ['Site and system information for the controller', 'The current operating status of the systems it controls',
              'Actions and recommendations, including work completed during the review']:
        w.bullet(b)
    w.pagebreak()

    # ---- 1 site information
    w.h1('Site and System Information')
    rows = [['Site', site], ['Controller', f"{station} – Niagara {meta.get('niagara') or '?'}"]]
    for k, label in (('supervisor', 'Supervisor'), ('groups', 'Terminal units'), ('devices', 'Field devices'), ('metering', 'Metering')):
        if summ.get(k):
            rows.append([label, summ[k]])
    rows += [['Review method', f"Remote review of the station configuration saved {meta.get('saved') or ''}"],
             ['Review date', meta.get('created', '')[:10]]]
    w.table(['Item', 'Detail'], rows, [4.5, 11.5])

    # ---- 2 status
    w.h1('System Status')
    if stat:
        w.para('The table below summarises how the systems are configured and operating.')
        w.table(['Item', 'Status', 'Result'],
                [[s['title'], s.get('clientText', ''), (s.get('result') or 'Satisfactory', AMBER if s.get('result') == 'See actions' else (GREEN if (s.get('result') or 'Satisfactory') == 'Satisfactory' else '3F6C88'), True)]
                 for s in stat], [4, 9.5, 2.5])
    if readings_item:
        h = w.h2('Readings at the Time of Review')
        p = w.para(readings_item.get('clientText', ''), keep=True)
        w.table(['Point', 'Reading'], [[r['point'], 'Sensor fault' if r.get('fault') else f"{_fmt(r['value'])} {r['units']}"]
                                       for r in readings_item.get('readings', [])], [8, 8])

    # ---- 3 actions
    w.pagebreak()
    w.h1('Actions and Recommendations')
    w.h2('Action Summary')
    pending = [a for a in acts if a not in done]
    if pending:
        w.para('The table below lists the actions arising from this review.')
        w.table(['Ref', 'Priority', 'Finding', 'Recommendation', 'Status', 'By'],
                [[f"A{n}", (a.get('priority', ''), PRI.get(a.get('priority'), '3F6C88'), True), a['title'], a.get('clientText', '').strip(),
                  a.get('status', 'Open'), 'Client' if a.get('status') == 'Client decision' else 'SSE'] for n, a in enumerate(pending, 1)],
                [1.1, 1.8, 3.6, 6.3, 2, 1.8])
    else:
        w.para('No actions arose from this review.')
    w.h2('Completed During Review')
    if done:
        w.para('The following items were corrected remotely during the review.')
        w.table(['Ref', 'Item', 'Action taken', 'Notes'], [[f"C{n}", a['title'], a.get('clientText', ''), ''] for n, a in enumerate(done, 1)],
                [1.1, 4.4, 6.5, 4])
    else:
        w.para('No items were corrected remotely during this review.')
        w.table(['Ref', 'Item', 'Action taken', 'Notes'], [['', '', '', ''], ['', '', '', '']], [1.1, 4.4, 6.5, 4])
    if reviewed:
        w.h2('Reviewed – No Action Required')
        w.table(['Item', 'Outcome'], [[r['title'], 'No action required. ' + (r.get('reportReason') or _plain_reason(r.get('disregardReason', '')))] for r in reviewed], [5.5, 10.5])

    # ---- 4 engineer notes
    w.h1('Engineer Notes')
    w.para('Space for notes on actions completed or follow-up visits.')
    w.table(['Date', 'Ref', 'Note', 'By'], [['', '', '', ''] for _ in range(4)], [2.5, 1.5, 9.5, 2.5])
    w.save(out_path)
    return out_path


def _plain_reason(r):
    """Strip the 'Rule:' / 'Decided on X:' prefixes so the client sees just the reason."""
    import re
    r = re.sub(r'^Rule:\s*', '', r or '')
    r = re.sub(r'^Decided on [^:]+:\s*', '', r)
    return r.strip()
